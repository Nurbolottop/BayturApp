import io
from datetime import timedelta

from django.core.management import call_command
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from apps.analytics import reports
from apps.analytics.models import AppEvent, DeviceLink, ExportJob
from apps.analytics.reports import ReportParams
from apps.analytics.tasks import run_export
from apps.cashback.services import confirm_request, create_request, credit_request
from apps.common.tokens import issue_access
from apps.loyalty.models import Operation, OperationKind
from apps.loyalty.services import manual_adjustment
from apps.members.models import Member
from apps.staff.models import StaffUser
from apps.staff.roles import Role


def kpi(data, key):
    return next(k for k in data['kpis'] if k['key'] == key)


def table(data, key):
    return next(t for t in data['tables'] if t['key'] == key)


class Dataset:
    """
    m1: заявка room-standard 14 000 сом наличными → кешбек 98 000; вторая заявка (pending) — повтор.
    m2: +100 000 корректировкой, заявка room-standard с оплатой 1 000 сом баллами → списано 100 000,
        деньгами 13 000, кешбек 91 000.
    mt: тестовый аккаунт с такой же заявкой — нигде не учитывается.
    """

    @classmethod
    def build(cls, tc):
        call_command('seed', stdout=io.StringIO())
        tc.cashier = StaffUser.objects.create_user('cashier@baytur.kg', 'x' * 12, full_name='Кассир',
                                                   role=Role.STAFF)
        now = timezone.now()
        tc.m1 = Member.objects.create(phone='+996700000001', first_name='A', first_device_id='d1')
        tc.m2 = Member.objects.create(phone='+996700000002', first_name='B')
        tc.mt = Member.objects.create(phone='+996700000009', first_name='T', is_test=True)

        def full(member, data):
            req, _ = create_request(member, data)
            confirm_request(req.pk, staff=tc.cashier, cash_received=True)
            return credit_request(req.pk)

        tc.r1 = full(tc.m1, {'itemId': 'room-standard', 'quantity': 1, 'method': 'cash'})
        create_request(tc.m1, {'itemId': 'room-standard', 'quantity': 1, 'method': 'cash'})  # повтор, pending
        manual_adjustment(tc.m2, 100_000, 'стартовые баллы', None)
        tc.r2 = full(tc.m2, {'itemId': 'room-standard', 'quantity': 1, 'method': 'cash', 'pointsSom': 1000})
        full(tc.mt, {'itemId': 'room-standard', 'quantity': 1, 'method': 'cash'})

        # события: d1 → m1 (регистрация), d2 — гость, d3 → m2, dt → тестовый
        hour_ago = now - timedelta(hours=1)
        ev = [('d1', None, 'app_open'), ('d1', None, 'login_prompt_shown'), ('d1', None, 'login_prompt_accepted'),
              ('d1', tc.m1, 'registration_completed'), ('d2', None, 'app_open'),
              ('d2', None, 'login_prompt_shown'), ('d3', None, 'app_open'), ('dt', tc.mt, 'app_open')]
        AppEvent.objects.bulk_create([AppEvent(device_id=d, member=m, name=n, at=hour_ago, platform='ios',
                                               language='ru') for d, m, n in ev])
        DeviceLink.objects.create(device_id='d1', member=tc.m1, platform='ios')
        DeviceLink.objects.create(device_id='d2', member=None, platform='ios')
        DeviceLink.objects.create(device_id='d3', member=tc.m2, platform='ios')
        DeviceLink.objects.create(device_id='dt', member=tc.mt, platform='ios')


class ReportsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        Dataset.build(cls)

    def params(self, **kw):
        return ReportParams.from_query({**kw})

    def test_dashboard_numbers(self):
        self.assertEqual(self.r1.cashback, 98_000)
        self.assertEqual((self.r2.money_som, self.r2.points, self.r2.cashback), (13_000, 100_000, 91_000))
        d = reports.dashboard(self.params())
        self.assertEqual(kpi(d, 'points_credited')['value'], 189_000)
        self.assertEqual(kpi(d, 'points_spent')['value'], 100_000)
        self.assertEqual(kpi(d, 'revenue')['value'], 27_000)
        self.assertEqual(kpi(d, 'requests')['value'], 3)
        self.assertEqual(kpi(d, 'requests_done')['value'], 2)
        self.assertEqual(kpi(d, 'new_members')['value'], 2)
        self.assertEqual(kpi(d, 'active_members')['value'], 2)
        self.assertEqual(kpi(d, 'new_members')['prev'], 0)
        # обязательство: 98 000 (m1) + 100 000 − 100 000 + 91 000 (m2)
        self.assertEqual(d['liability']['points'], 189_000)
        self.assertEqual(d['liability']['som'], 1890)
        cat = table(d, 'by_category')['rows']
        self.assertEqual([(r['category_id'], r['requests'], r['revenue']) for r in cat], [('rooms', 2, 27_000)])
        method = table(d, 'by_method')['rows']
        self.assertEqual(method[0]['method_id'], 'cash')
        staff = table(d, 'staff')['rows']
        self.assertEqual((staff[0]['staff'], staff[0]['confirmed']), ('Кассир', 2))
        tiers = {r['tier_id']: r['members'] for r in table(d, 'tiers')['rows']}
        self.assertEqual(sum(tiers.values()), 2)
        series = next(s for s in d['series'] if s['key'] == 'points_credited')
        self.assertEqual(len(series['points']), 30)
        self.assertEqual(sum(pt['y'] for pt in series['points']), 189_000)

    def test_test_accounts_excluded(self):
        # тестовый аккаунт имеет начисление, но его нет ни в одном отчёте
        self.assertTrue(Operation.objects.filter(member=self.mt, kind=OperationKind.CASHBACK).exists())
        m = reports.money(self.params())
        self.assertEqual(kpi(m, 'revenue')['value'], 27_000)
        self.assertEqual(kpi(m, 'requests')['value'], 2)
        p = reports.points(self.params())
        self.assertEqual(kpi(p, 'credited')['value'], 189_000)
        self.assertEqual(kpi(p, 'adjust_plus')['value'], 100_000)
        self.assertEqual(reports.members(self.params())['kpis'][0]['value'], 2)
        self.assertEqual(reports.liability(timezone.localdate())['points'], 189_000)

    def test_filters(self):
        d = reports.dashboard(self.params(category='spa'))
        self.assertEqual(kpi(d, 'revenue')['value'], 0)
        d = reports.dashboard(self.params(outlet='reception', groupBy='month'))
        self.assertEqual(kpi(d, 'revenue')['value'], 27_000)
        self.assertEqual(reports.dashboard(self.params(platform='android'))['kpis'][2]['value'], 0)

    def test_liability_on_past_date(self):
        today = timezone.localdate()
        Operation.objects.filter(member=self.m1).update(at=timezone.now() - timedelta(days=10))
        self.assertEqual(reports.liability(today - timedelta(days=5))['points'], 98_000)
        self.assertEqual(reports.liability(today - timedelta(days=5))['som'], 980)
        self.assertEqual(reports.liability(today - timedelta(days=15))['points'], 0)
        self.assertEqual(reports.liability(today)['points'], 189_000)

    def test_funnel_counts(self):
        f = reports.funnel(self.params())
        counts = {s['key']: s['count'] for s in f['steps']}
        self.assertEqual(counts, {'opened': 3, 'registered': 2, 'first_req': 2, 'first_credit': 2, 'second_req': 1})
        reg = next(s for s in f['steps'] if s['key'] == 'registered')
        self.assertEqual(reg['conversion'], 0.6667)
        self.assertEqual(f['login'], {'shown': 2, 'accepted': 1, 'conversion': 0.5})

    def test_other_blocks_render(self):
        for name in ('cohorts', 'catalog', 'content', 'operations', 'complaints'):
            data = reports.run(name, {})
            self.assertIn('kpis', data)
            self.assertIn('tables', data)
        ops = reports.operations(self.params())
        self.assertEqual(kpi(ops, 'confirmed')['value'], 2)
        self.assertIsNotNone(kpi(ops, 'median_confirm_minutes')['value'])
        cat = reports.catalog(self.params())
        self.assertEqual(table(cat, 'top_items')['rows'][0]['item_id'], 'room-standard')


@override_settings(ROOT_URLCONF='apps.analytics.tests.urls')
class AdminApiTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        Dataset.build(cls)
        cls.owner = StaffUser.objects.create_user('owner@baytur.kg', 'x' * 12, full_name='O', role=Role.OWNER)
        # администратор кассы — без аналитики; второй директор — для проверки чужих выгрузок
        cls.editor = StaffUser.objects.create_user(None, None, full_name='E', role=Role.STAFF, phone='+996700000077')
        cls.accountant = StaffUser.objects.create_user('acc@baytur.kg', 'x' * 12, full_name='A', role=Role.OWNER)
        cls.other = StaffUser.objects.create_user('other@baytur.kg', 'x' * 12, full_name='B', role=Role.OWNER)

    def get(self, url, user):
        return self.client.get(url, HTTP_AUTHORIZATION=f'Bearer {issue_access("staff", user.pk)}')

    def post(self, url, user, data):
        return self.client.post(url, data, content_type='application/json',
                                HTTP_AUTHORIZATION=f'Bearer {issue_access("staff", user.pk)}')

    def test_permissions(self):
        self.assertEqual(self.client.get('/api/v1/admin/analytics/money').status_code, 401)
        self.assertEqual(self.get('/api/v1/admin/analytics/money', self.editor).status_code, 403)
        self.assertEqual(self.get('/api/v1/admin/reports/dashboard', self.editor).status_code, 403)
        self.assertEqual(self.get('/api/v1/admin/analytics/content', self.editor).status_code, 403)
        self.assertEqual(self.get('/api/v1/admin/analytics/money', self.accountant).status_code, 200)
        self.assertEqual(self.get('/api/v1/admin/analytics/funnel', self.accountant).status_code, 200)
        r = self.get('/api/v1/admin/reports/dashboard?groupBy=week', self.owner)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(kpi(r.json(), 'revenue')['value'], 27_000)
        r = self.get(f'/api/v1/admin/analytics/liability?date={timezone.localdate()}', self.accountant)
        self.assertEqual(r.json()['points'], 189_000)
        self.assertEqual(self.get('/api/v1/admin/analytics/complaints', self.owner).status_code, 200)
        # экспорт — только тех отчётов, которые пользователь видит
        r = self.post('/api/v1/admin/analytics/export', self.editor, {'report': 'money', 'format': 'xlsx'})
        self.assertEqual(r.status_code, 403)

    def test_export_xlsx(self):
        with self.captureOnCommitCallbacks(execute=False):
            r = self.post('/api/v1/admin/analytics/export', self.accountant,
                          {'report': 'dashboard', 'params': {'groupBy': 'week'}, 'format': 'xlsx'})
        self.assertEqual(r.status_code, 202, r.content)
        job = ExportJob.objects.get(pk=r.json()['id'])
        self.assertEqual(job.params['groupBy'], 'week')
        self.assertEqual(run_export(job.pk), 'done')
        job.refresh_from_db()
        try:
            self.assertTrue(job.file.name.endswith('.xlsx'))
            with job.file.open('rb') as fh:
                wb = load_workbook(io.BytesIO(fh.read()))
            self.assertIn('Показатели', wb.sheetnames)
            values = [row[1] for row in wb['Показатели'].iter_rows(min_row=2, values_only=True)]
            self.assertIn(27_000, values)
            listing = self.get('/api/v1/admin/analytics/exports', self.accountant).json()['items']
            self.assertEqual(listing[0]['status'], 'done')
            dl = self.get(listing[0]['downloadUrl'], self.accountant)
            self.assertEqual(dl.status_code, 200)
            self.assertTrue(b''.join(dl.streaming_content).startswith(b'PK'))
            # администратору кассы выгрузки недоступны
            self.assertIn(self.get(listing[0]['downloadUrl'], self.editor).status_code, (403, 404))
        finally:
            job.file.delete(save=False)

    def test_export_csv(self):
        job = ExportJob.objects.create(report='points', params={}, fmt='csv', created_by=self.owner)
        run_export(job.pk)
        job.refresh_from_db()
        try:
            self.assertEqual(job.status, 'done', job.error)
            with job.file.open('rb') as fh:
                text = fh.read().decode('utf-8-sig')
            self.assertIn('# Показатели', text)
        finally:
            job.file.delete(save=False)


class SmokeTests(TestCase):
    """Все блоки на данных с событиями контента, рассылкой и обращениями — и все JSON-сериализуемы."""

    @classmethod
    def setUpTestData(cls):
        Dataset.build(cls)
        from apps.complaints.models import Complaint
        from apps.notifications.models import Campaign, Notification
        now = timezone.now()
        camp = Campaign.objects.create(title={'ru': 'Осень'}, body={'ru': '...'}, status='sent', sent_at=now)
        Notification.objects.create(member=cls.m1, kind='campaign', title='t', body='b', campaign=camp,
                                    push_status='sent', opened_at=now, created_at=now - timedelta(hours=3))
        Notification.objects.create(member=cls.m2, kind='campaign', title='t', body='b', campaign=camp,
                                    push_status='failed', created_at=now - timedelta(hours=3))
        Complaint.objects.create(seq=1, member=cls.m1, category_id='service', outlet_id='reception', rating=4,
                                 first_reply_at=now, closed_at=now, due_at=now + timedelta(hours=1),
                                 created_at=now - timedelta(hours=2))
        Complaint.objects.create(seq=2, member=cls.mt, category_id='service', rating=1)
        at = now - timedelta(minutes=30)
        AppEvent.objects.bulk_create([
            AppEvent(device_id='d1', name='article_view', at=at, props={'articleId': 'spa-bochka-promo'}),
            AppEvent(device_id='d1', name='promo_click', at=at, props={'promoId': 'spa-bochka-promo'}),
            AppEvent(device_id='d1', name='story_view', at=at, props={'category': 'spa', 'slide': 0}),
            AppEvent(device_id='d1', name='story_view', at=at, props={'category': 'spa', 'slide': 1}),
            AppEvent(device_id='d3', name='story_view', at=at, props={'category': 'spa', 'slide': 0}),
            AppEvent(device_id='d1', name='cashback_flow_started', at=at, props={'source': 'story'}),
            AppEvent(device_id='d1', name='cashback_flow_submitted', at=at, props={'source': 'story'}),
        ])

    def test_all_reports_serializable(self):
        import json
        for name in reports.REPORTS:
            for q in ({}, {'groupBy': 'week', 'tier': 'bronze'}, {'groupBy': 'month', 'language': 'ru',
                                                                   'platform': 'ios', 'compare': '0'}):
                json.dumps(reports.run(name, q))

    def test_content_campaigns_complaints(self):
        c = reports.content(ReportParams.from_query({}))
        self.assertEqual(kpi(c, 'article_views')['value'], 1)
        self.assertEqual(kpi(c, 'cta_requests')['value'], 1)
        camp = table(c, 'campaigns')['rows'][0]
        self.assertEqual((camp['sent'], camp['delivered'], camp['opened'], camp['requests_7d']), (2, 1, 1, 2))
        stories = {r['category_id']: r for r in table(c, 'stories')['rows']}
        self.assertEqual(stories['spa']['viewers'], 2)
        comp = reports.complaints_report(ReportParams.from_query({}))
        self.assertEqual(kpi(comp, 'complaints')['value'], 1)  # тестовый не учтён
        self.assertEqual(kpi(comp, 'avg_rating')['value'], 4)
        self.assertEqual(kpi(comp, 'avg_first_response_hours')['value'], 2)
        cohorts = reports.cohorts(ReportParams.from_query({}))
        self.assertEqual(sum(r['size'] for r in cohorts['cohorts']['rows']), 2)
