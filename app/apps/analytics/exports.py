"""
Выгрузка отчётов в XLSX / CSV (§7.2, §7.5): фоновая задача пишет файл в приватное хранилище,
ссылка на скачивание появляется в админке (GET /admin/analytics/exports).
"""
import csv
import io
import logging
import re

from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone

from . import reports
from .models import ExportJob

log = logging.getLogger(__name__)

FORMATS = ('xlsx', 'csv')
CONTENT_TYPES = {
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'csv': 'text/csv; charset=utf-8',
}


def _cell(v):
    if isinstance(v, (dict, list)):
        return ', '.join(f'{x}' for x in (v.values() if isinstance(v, dict) else v))
    return v


def sections(report, data):
    """Отчёт → список секций (title, header, rows) — общий формат для XLSX и CSV."""
    out = []
    period = data.get('period') or {}
    meta = [['Отчёт', reports.REPORT_TITLES.get(report, report)],
            ['Сформирован', timezone.localtime().strftime('%Y-%m-%d %H:%M')]]
    if period:
        meta += [['Период', f"{period.get('from')} — {period.get('to')}"],
                 ['Группировка', period.get('groupBy')]]
        if period.get('prev'):
            meta.append(['Сравнение с', f"{period['prev']['from']} — {period['prev']['to']}"])
        for k, v in (period.get('filters') or {}).items():
            meta.append([f'Фильтр: {k}', v])
    elif data.get('date'):
        meta.append(['На дату', data['date']])
    out.append(('Параметры', ['Параметр', 'Значение'], meta))

    kpis = data.get('kpis') or []
    if kpis:
        out.append(('Показатели', ['Показатель', 'Значение', 'Прошлый период', 'Ед.'],
                    [[k['label'], k['value'], k.get('prev'), k.get('unit')] for k in kpis]))

    series = data.get('series') or []
    if series:
        xs = []
        for s in series:
            for pt in s['points']:
                if pt['x'] not in xs:
                    xs.append(pt['x'])
        by = [{pt['x']: pt['y'] for pt in s['points']} for s in series]
        out.append(('Динамика', ['Дата'] + [s['label'] for s in series],
                    [[x] + [b.get(x) for b in by] for x in xs]))

    for t in data.get('tables') or []:
        cols = t.get('columns') or []
        if not cols and t.get('rows'):
            cols = [{'key': k, 'label': k} for k in t['rows'][0].keys()]
        out.append((t.get('label') or t.get('key'), [c['label'] for c in cols],
                    [[_cell(r.get(c['key'])) for c in cols] for r in t.get('rows') or []]))
    return out


def _sheet_title(title, used):
    base = re.sub(r'[\[\]\*\?/\\:]', ' ', str(title or 'Лист')).strip()[:31] or 'Лист'
    name, i = base, 2
    while name in used:
        suffix = f' ({i})'
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


def to_xlsx(secs) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for title, header, rows in secs:
        ws = wb.create_sheet(_sheet_title(title, used))
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        widths = [len(str(h)) for h in header]
        for r in rows:
            ws.append(list(r))
            for i, v in enumerate(r):
                if i < len(widths):
                    widths[i] = max(widths[i], min(60, len(str(v)) if v is not None else 0))
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, w + 2)
        ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(secs) -> bytes:
    """Один CSV: секции подряд, у каждой строка-заголовок с названием; разделитель «;» (Excel в ru-локали)."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=';')
    for i, (title, header, rows) in enumerate(secs):
        if i:
            w.writerow([])
        w.writerow([f'# {title}'])
        w.writerow(header)
        for r in rows:
            w.writerow(['' if v is None else v for v in r])
    return ('﻿' + buf.getvalue()).encode('utf-8')


def render(report, params, fmt='xlsx'):
    """→ (filename, bytes). params — dict в query-формате (from, to, groupBy, …) или {'date': …} для liability."""
    if report not in reports.REPORTS:
        raise ValueError(f'unknown report {report}')
    if fmt not in FORMATS:
        raise ValueError(f'unknown format {fmt}')
    params = dict(params or {})
    if report == 'liability' and params.get('date') and not params.get('to'):
        params['to'] = params['date']
    data = reports.run(report, params)
    secs = sections(report, data)
    payload = to_xlsx(secs) if fmt == 'xlsx' else to_csv(secs)
    period = data.get('period') or {}
    span = f"{period.get('from')}_{period.get('to')}" if period else (data.get('date') or '')
    return f'{report}_{span}.{fmt}'.replace(':', '-'), payload


def create_job(user, report, params, fmt='xlsx'):
    job = ExportJob.objects.create(report=report, params=params or {}, fmt=fmt, created_by=user)
    transaction.on_commit(lambda: enqueue(job.pk))
    return job


def enqueue(job_id):
    from .tasks import run_export
    try:
        run_export.delay(job_id)
    except Exception:  # брокер недоступен — задачу можно перезапустить из админки
        log.warning('export enqueue failed for job %s', job_id)


def run_job(job_id):
    job = ExportJob.objects.filter(pk=job_id).first()
    if job is None or job.status == 'done':
        return job
    try:
        name, payload = render(job.report, job.params, job.fmt)
        job.file.save(name, ContentFile(payload), save=False)
        job.status = 'done'
        job.error = ''
    except Exception as exc:  # noqa: BLE001 — ошибка видна в админке
        log.exception('export job %s failed', job_id)
        job.status = 'failed'
        job.error = f'{type(exc).__name__}: {exc}'[:2000]
    job.finished_at = timezone.now()
    job.save(update_fields=['file', 'status', 'error', 'finished_at'])
    return job
