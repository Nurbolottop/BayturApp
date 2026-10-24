"""
Сверка оплат баллами в R-Keeper со списаниями BAYTUR.

Кассир ресторана списывает баллы в приложении и получает код из 6 цифр (CashbackRequest.pos_code), затем в
R-Keeper закрывает чек оплатой «Баллы BAYTUR» и вписывает код в комментарий. Раз в день менеджер загружает
выгрузку оплат «Баллы» из R-Keeper (XLSX или CSV) — каждая строка сверяется по коду и сумме:

- red    — оплата баллами в R-Keeper, а списания у нас нет (нет кода / код не найден): подозрительно;
- yellow — код есть, но сумма в R-Keeper больше списанной, списание отменено или не подтверждено, код повторяется;
- orange — баллы у нас списаны, а в R-Keeper такой оплаты нет (продажа не пробита);
- ok     — совпало.

Колонки выгрузки находятся по заголовкам («Дата», «Сумма», «Комментарий», «Кассир», «Чек»); код ищется
в комментарии, а если такой колонки нет — в любой текстовой ячейке строки.
"""
import csv
import io
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from django.utils import timezone

from apps.common.errors import ApiError

CODE_RE = re.compile(r'(?<!\d)\d{6}(?!\d)')
COLUMNS = {  # поле → слова в заголовке (по порядку приоритета)
    'amount': ('сумма', 'amount', 'sum'),
    'at': ('дата', 'время', 'закрыт', 'date', 'time'),
    'comment': ('коммент', 'примеч', 'код', 'comment', 'note'),
    'cashier': ('кассир', 'официант', 'сотрудник', 'менеджер', 'cashier', 'waiter'),
    'check': ('чек', 'заказ', 'check', 'order'),
}
MAX_ROWS = 20000
DONE = ('confirmed', 'credited')


# ---------------------------------------------------------------- чтение файла

def _table(f, name):
    """Файл → список строк (списки ячеек). XLSX — первый лист; CSV — разделитель ; или , (UTF-8 / Windows-1251)."""
    data = f.read()
    if name.lower().endswith(('.xlsx', '.xlsm')) or data[:2] == b'PK':
        from openpyxl import load_workbook
        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception:
            raise ApiError('validation_error', 400, message='Не удалось открыть Excel-файл')
        return [list(r) for r, _ in zip(ws.iter_rows(values_only=True), range(MAX_ROWS))]
    if name.lower().endswith('.xls'):
        raise ApiError('validation_error', 400, message='Старый формат .xls: сохраните выгрузку как .xlsx или .csv')
    for enc in ('utf-8-sig', 'cp1251'):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ApiError('validation_error', 400, message='Не удалось прочитать файл')
    dialect = csv.Sniffer().sniff(text[:4000], delimiters=';,\t') if text.strip() else csv.excel
    return [row for row, _ in zip(csv.reader(io.StringIO(text), dialect), range(MAX_ROWS))]


def _header(rows):
    """Строка заголовков (среди первых 30) и номера колонок по полям."""
    for i, row in enumerate(rows[:30]):
        cells = [str(c or '').strip().lower() for c in row]
        cols = {}
        for field, words in COLUMNS.items():
            for word in words:
                j = next((j for j, c in enumerate(cells) if word in c and j not in cols.values()), None)
                if j is not None:
                    cols[field] = j
                    break
        if 'amount' in cols and len(cols) >= 2:
            return i, cols
    raise ApiError('validation_error', 400,
                   message='Не нашёл в файле колонку «Сумма». Нужна выгрузка оплат из R-Keeper с заголовками')


def _amount(v):
    if isinstance(v, (int, float, Decimal)):
        return int(Decimal(str(v)).quantize(Decimal('1')))
    s = re.sub(r'[^\d,.\-]', '', str(v or '')).replace(',', '.')
    if s.count('.') > 1:  # 12.500.00 — разделители тысяч
        s = s.replace('.', '', s.count('.') - 1)
    try:
        return int(Decimal(s).quantize(Decimal('1')))
    except InvalidOperation:
        return None


def _at(v):
    if isinstance(v, datetime):
        return timezone.make_aware(v) if timezone.is_naive(v) else v
    if isinstance(v, date):
        return timezone.make_aware(datetime.combine(v, time()))
    s = str(v or '').strip()
    for fmt in ('%d.%m.%Y %H:%M:%S', '%d.%m.%Y %H:%M', '%d.%m.%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M',
                '%Y-%m-%d', '%d.%m.%y %H:%M', '%d/%m/%Y %H:%M'):
        try:
            return timezone.make_aware(datetime.strptime(s, fmt))
        except ValueError:
            continue
    return None


def read_rows(f, name):
    """Выгрузка R-Keeper → [{n, at, amount, code, cashier, check, text}] — только строки с суммой, без «Итого»."""
    rows = _table(f, name)
    head, cols = _header(rows)
    out = []
    for n, row in enumerate(rows[head + 1:], start=head + 2):
        cell = lambda field: row[cols[field]] if field in cols and cols[field] < len(row) else None  # noqa: E731
        texts = [str(c).strip() for c in row if c not in (None, '')]
        if not texts or any('итог' in t.lower() or 'всего' in t.lower() for t in texts):
            continue
        amount = _amount(cell('amount'))
        if amount is None or amount <= 0:
            continue
        if 'comment' in cols:
            m = CODE_RE.search(str(cell('comment') or ''))
        else:  # кода в отдельной колонке нет — ищем в любой текстовой ячейке, кроме суммы, даты и номера чека
            skip = {cols.get(k) for k in ('amount', 'at', 'check')}
            m = next((CODE_RE.search(str(c)) for j, c in enumerate(row)
                      if j not in skip and isinstance(c, str) and CODE_RE.search(c)), None)
        at = _at(cell('at'))
        out.append({'n': n, 'at': at.isoformat() if at else None, 'amount': amount, 'code': m.group() if m else None,
                    'cashier': str(cell('cashier') or '').strip(), 'check': str(cell('check') or '').strip(),
                    'text': ' · '.join(texts)[:300]})
    if not out:
        raise ApiError('validation_error', 400, message='В файле нет строк с суммой оплаты')
    return out


# ---------------------------------------------------------------- сверка

def _req_info(req):
    return {'id': req.pk, 'code': req.pos_code, 'total': req.total, 'status': req.status, 'at': req.created_at.isoformat(),
            'outlet': req.outlet_id, 'member': req.member.member_id,
            'staff': getattr(req.created_by or req.confirmed_by, 'full_name', '') or ''}


def reconcile(rows, requests_qs, period_from, period_to):
    """
    rows — из read_rows; requests_qs — заявки, которые видит сотрудник (scoped_requests);
    period_from/period_to — даты (включительно), за которые выгрузка: по ним ищутся списания без оплаты в R-Keeper.
    → {'rows': [...со status/reason/request], 'missing': [...], 'summary': {...}, 'byCashier': [...]}
    """
    from .models import CashbackRequest
    codes = {r['code'] for r in rows if r['code']}
    by_code = {req.pos_code: req for req in CashbackRequest.objects.filter(pos_code__in=codes)
               .select_related('member', 'created_by', 'confirmed_by')}
    visible = set(requests_qs.filter(pos_code__in=codes).values_list('pk', flat=True))
    seen = {}
    for r in rows:
        req = by_code.get(r['code']) if r['code'] else None
        if req is not None and req.pk not in visible:
            req = None  # чужой отдел / точка — для этого сотрудника кода нет
        if not r['code']:
            r['status'], r['reason'] = 'red', 'Нет кода BAYTUR в комментарии'
        elif req is None:
            r['status'], r['reason'] = 'red', 'Код не найден: списания баллов не было'
        elif r['code'] in seen:
            r['status'], r['reason'] = 'yellow', f'Код уже был в строке {seen[r["code"]]}'
        elif req.status not in DONE:
            r['status'], r['reason'] = 'yellow', ('Списание отменено' if req.status in ('cancelled', 'rejected')
                                                  else 'Клиент не подтвердил списание')
        elif r['amount'] > req.total:
            r['status'], r['reason'] = 'yellow', f'В R-Keeper {r["amount"]} сом, списано на {req.total} сом'
        else:
            r['status'], r['reason'] = 'ok', ''
        if r['code'] and r['code'] not in seen:
            seen[r['code']] = r['n']
        r['request'] = _req_info(req) if req is not None else None

    start = timezone.make_aware(datetime.combine(period_from, time.min))
    end = timezone.make_aware(datetime.combine(period_to, time.max))
    matched = {r['request']['id'] for r in rows if r['request']}
    missing = [_req_info(req) for req in requests_qs.filter(
        pos_code__isnull=False, status__in=DONE, outlet__receipt_required=True, created_at__range=(start, end))
        .exclude(pk__in=matched).select_related('member', 'created_by', 'confirmed_by').order_by('created_at')]

    by_cashier = {}
    for r in rows:
        if r['status'] == 'ok':
            continue
        row = by_cashier.setdefault(r['cashier'] or '—', {'name': r['cashier'] or '—', 'red': 0, 'yellow': 0,
                                                          'sum': 0})
        row[r['status']] += 1
        row['sum'] += r['amount']
    summary = {k: sum(1 for r in rows if r['status'] == k) for k in ('ok', 'red', 'yellow')}
    summary.update(orange=len(missing), rows=len(rows),
                   redSum=sum(r['amount'] for r in rows if r['status'] == 'red'))
    return {'rows': rows, 'missing': missing, 'summary': summary,
            'byCashier': sorted(by_cashier.values(), key=lambda x: (-x['red'], -x['sum']))}


def period_of(rows):
    """Даты первой и последней строки выгрузки (локальные), если в файле есть дата."""
    days = [timezone.localtime(datetime.fromisoformat(r['at'])).date() for r in rows if r['at']]
    return (min(days), max(days)) if days else (None, None)
