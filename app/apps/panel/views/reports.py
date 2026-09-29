"""
Дашборд и «Аналитика» (ТЗ §7.2, §7.5). Данные — из apps.analytics.reports (пишется параллельно):
каждый отчёт → {kpis, series, tables}. Пока модуля нет — страница «Отчёты ещё не готовы».
"""
import csv
import io
from datetime import date

from django.contrib import messages
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from django.views.decorators.http import require_POST

from apps.catalog.models import Category, Outlet
from apps.loyalty.models import TierId

from ..access import forbidden, panel_view
from ..sections import ANALYTICS_BLOCKS


def _reports():
    try:
        from apps.analytics import reports
    except ImportError:
        return None
    return reports


def _run(reports, fn_name, request):
    fn = getattr(reports, fn_name, None)
    if fn is None:
        return None, 'Отчёт ещё не готов'
    try:
        params = reports.ReportParams.from_query(request.GET)
        return normalize(fn(params)), None
    except Exception as e:  # отчёт не должен ронять страницу
        import logging
        logging.getLogger(__name__).exception('report %s failed', fn_name)
        return None, f'Ошибка отчёта: {e.__class__.__name__}'


def normalize(data):
    data = dict(data or {})
    data.setdefault('kpis', [])
    data.setdefault('series', [])
    data.setdefault('tables', [])
    return data


def filters_ctx(request):
    return {
        'f': request.GET,
        'categories': Category.objects.all(),
        'outlets': Outlet.objects.all(),
        'tiers': TierId.choices,
        'group_by': [('day', 'Дни'), ('week', 'Недели'), ('month', 'Месяцы')],
        'platforms': [('ios', 'iOS'), ('android', 'Android')],
        'languages': [('ru', 'RU'), ('ky', 'KY'), ('en', 'EN')],
    }


@panel_view('dashboard')
def dashboard(request):
    reports = _reports()
    ctx = filters_ctx(request)
    if reports is None:
        ctx['not_ready'] = True
    else:
        ctx['report'], ctx['error'] = _run(reports, 'dashboard', request)
    return render(request, 'panel/reports/dashboard.html', ctx)


def allowed_blocks(user):
    return [b for b in ANALYTICS_BLOCKS if user.can(b[2])]


@panel_view('analytics')
def analytics(request, block=None):
    blocks = allowed_blocks(request.user)
    if not blocks:
        return forbidden(request)
    keys = {b[0]: b for b in blocks}
    if block is None:
        return redirect('panel:analytics-block', block=blocks[0][0])
    if block not in keys:
        if block in {b[0] for b in ANALYTICS_BLOCKS}:
            return forbidden(request, 'Этот блок аналитики недоступен вашей роли')
        raise Http404
    reports = _reports()
    ctx = filters_ctx(request)
    ctx.update({'blocks': blocks, 'cur': keys[block], 'can_export': request.user.can('analytics.export')})
    if reports is None:
        ctx['not_ready'] = True
    else:
        ctx['report'], ctx['error'] = _run(reports, keys[block][3], request)
        if block == 'points' and hasattr(reports, 'liability'):
            try:
                on = date.fromisoformat(request.GET.get('liability_date') or date.today().isoformat())
            except ValueError:
                on = date.today()
            try:
                ctx['liability'] = normalize(reports.liability(on))
            except Exception:
                ctx['liability'] = None
            ctx['liability_date'] = on
    return render(request, 'panel/reports/analytics.html', ctx)


@require_POST
@panel_view(perm='analytics.export')
def analytics_export(request, block):
    """
    Выгрузка XLSX/CSV: apps.analytics.exports.render (тот же формат, что у фоновой выгрузки Admin-API);
    если модуля нет — CSV/XLSX из таблиц отчёта.
    """
    keys = {b[0]: b for b in allowed_blocks(request.user)}
    if block not in keys and block != 'dashboard':
        return forbidden(request)
    fmt = 'xlsx' if request.POST.get('format') == 'xlsx' else 'csv'
    reports = _reports()
    if reports is None:
        messages.error(request, 'Отчёты ещё не готовы')
        return redirect(request.POST.get('next') or 'panel:analytics')
    # ключ отчёта в apps.analytics.reports.REPORTS (обращения там — 'complaints')
    report_key = block
    params = {k: v for k, v in request.POST.items() if k not in ('csrfmiddlewaretoken', 'format', 'next')}
    try:
        from apps.analytics import exports
    except ImportError:
        exports = None
    if exports is not None and hasattr(exports, 'render'):
        try:
            filename, payload = exports.render(report_key, params, fmt)
        except (ValueError, KeyError):
            filename = None
        if filename:
            from apps.common.audit import audit
            audit(request, 'analytics.export', None, object_type='analytics', object_id=report_key,
                  after={'format': fmt, 'params': params})
            ctype = ('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if fmt == 'xlsx'
                     else 'text/csv; charset=utf-8')
            resp = HttpResponse(payload, content_type=ctype)
            resp['Content-Disposition'] = f'attachment; filename="{filename}"'
            return resp
    fn_name = 'dashboard' if block == 'dashboard' else keys[block][3]
    from django.http import QueryDict
    q = QueryDict(mutable=True)
    for k, v in request.POST.lists():
        if k not in ('csrfmiddlewaretoken', 'format', 'next'):
            q.setlist(k, v)
    request.GET = q
    data, error = _run(reports, fn_name, request)
    if data is None:
        messages.error(request, error or 'Не удалось выгрузить')
        return redirect(request.POST.get('next') or 'panel:analytics')
    name = f'baytur-{block}-{date.today().isoformat()}'
    if fmt == 'xlsx':
        return _xlsx(data, name)
    return _csv(data, name)


def _table_rows(t):
    cols = t.get('columns') or []
    yield [c.get('label') or c.get('key') for c in cols]
    for row in t.get('rows') or []:
        if isinstance(row, dict):
            yield [_plain(row.get(c.get('key'))) for c in cols]
        else:
            yield [_plain(v) for v in row]


def _plain(v):
    if isinstance(v, dict):
        return v.get('ru') or ''
    return v


def _kpi_rows(data):
    yield ['Показатель', 'Значение', 'Прошлый период', 'Ед.']
    for k in data['kpis']:
        yield [k.get('label'), k.get('value'), k.get('prev'), k.get('unit')]


def _csv(data, name):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=';')
    for row in _kpi_rows(data):
        w.writerow(row)
    for t in data['tables']:
        w.writerow([])
        w.writerow([t.get('label')])
        for row in _table_rows(t):
            w.writerow(row)
    for s in data['series']:
        w.writerow([])
        w.writerow([s.get('label')])
        for p in s.get('points') or []:
            w.writerow([p.get('x'), p.get('y')])
    resp = HttpResponse('﻿' + buf.getvalue(), content_type='text/csv; charset=utf-8')
    resp['Content-Disposition'] = f'attachment; filename="{name}.csv"'
    return resp


def _xlsx(data, name):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'KPI'
    for row in _kpi_rows(data):
        ws.append(row)
    for i, t in enumerate(data['tables']):
        sheet = wb.create_sheet((t.get('label') or f'Таблица {i + 1}')[:31].replace('/', '-'))
        for row in _table_rows(t):
            sheet.append([v if isinstance(v, (int, float, str)) or v is None else str(v) for v in row])
    for i, s in enumerate(data['series']):
        sheet = wb.create_sheet((s.get('label') or f'Ряд {i + 1}')[:31].replace('/', '-'))
        sheet.append(['x', 'y'])
        for p in s.get('points') or []:
            sheet.append([str(p.get('x')), p.get('y')])
    out = io.BytesIO()
    wb.save(out)
    resp = HttpResponse(out.getvalue(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{name}.xlsx"'
    return resp
